# report_generator.py
# -*- coding: utf-8 -*-
"""
Модуль генерации полного отчёта в формате DOCX.
"""

from datetime import datetime
from io import BytesIO

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


def _add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def _add_paragraph(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    return p


def _add_table_from_dict(doc, data, col1_name="Параметр", col2_name="Значение"):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    hdr[0].text = col1_name
    hdr[1].text = col2_name
    for cell in hdr:
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
    for key, value in data.items():
        cells = table.add_row().cells
        cells[0].text = str(key)
        cells[1].text = str(value)
    doc.add_paragraph()


def _add_table_from_list(doc, rows, headers):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = str(h)
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    doc.add_paragraph()


def _add_formula_block(doc, formula_text, description=""):
    if description:
        p = doc.add_paragraph()
        p.add_run(description).italic = True
    p = doc.add_paragraph()
    run = p.add_run(formula_text)
    run.font.name = "Consolas"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    doc.add_paragraph()


def generate_report(results, road_params, iri_forecast=None):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    # Титульный лист
    title = doc.add_heading("ОТЧЁТ", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("об оценке остаточного ресурса дорожной конструкции")
    run.bold = True
    run.font.size = Pt(14)
    doc.add_paragraph()
    doc.add_paragraph()
    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    date_p.add_run(f"Дата составления: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
    doc.add_paragraph()

    # 1. Исходные данные
    _add_heading(doc, "1. Исходные данные", level=1)
    _add_heading(doc, "1.1. Параметры дороги", level=2)
    _add_table_from_dict(doc, {
        "Категория дороги": road_params.get("road_category", "—"),
        "Расчётная нагрузка, кН": road_params.get("design_load", "—"),
        "Расчётный срок службы, лет": road_params.get("design_life", "—"),
        "Число полос движения": road_params.get("num_lanes", "—"),
        "Годовой прирост интенсивности, %": f"{road_params.get('q_growth_percent', 0):.1f}",
        "Расчётное число дней в году (Трдг)": road_params.get("trdg", "—"),
        "Фактический срок эксплуатации, лет": road_params.get("years_in_operation", "—"),
    })
    _add_heading(doc, "1.2. Текущее состояние", level=2)
    _add_table_from_dict(doc, {
        "Текущий IRI, м/км": f"{road_params.get('current_iri', 0):.2f}",
        "Критический порог IRI, м/км": f"{road_params.get('iri_limit', 0):.2f}",
    })
    _add_heading(doc, "1.3. Параметры прочности", level=2)
    _add_table_from_dict(doc, {
        "Тип дорожной одежды": road_params.get("road_type", "—"),
        "Категория для Kпр": road_params.get("kpr_category", "—"),
        "Kпр_тр (требуемый)": f"{road_params.get('kpr_tr', 0):.3f}",
        "Kпр_факт (фактический)": f"{road_params.get('kpr_fact', 0):.3f}",
    })

    # 2. Интенсивность
    doc.add_page_break()
    _add_heading(doc, "2. Интенсивность движения", level=1)
    traffic = results.get("traffic_by_type", {})
    if traffic:
        rows = [[v, f"{n:,}"] for v, n in traffic.items()]
        rows.append(["ИТОГО", f"{sum(traffic.values()):,}"])
        _add_table_from_list(doc, rows, ["Тип ТС", "Интенсивность, авт/сут"])
    else:
        _add_paragraph(doc, "Данные отсутствуют.", italic=True)

    # 3. Sm
    doc.add_page_break()
    _add_heading(doc, "3. Коэффициенты приведения Sm", level=1)
    _add_paragraph(doc, f"Метод: {results.get('sm_method', '—')}", bold=True)
    sm_by_type = results.get("sm_by_type", {})
    sm_info = results.get("sm_info_by_type", {})
    method_labels = {
        "table": "Таблица 6.1", "formula": "Формула 6.8",
        "table_fallback": "Таблица 6.1 (fallback)", "mixed": "Смешанный",
    }
    if sm_by_type:
        rows = []
        for v, sm in sorted(sm_by_type.items(), key=lambda x: -x[1]):
            method = sm_info.get(v, {}).get("method", "—")
            rows.append([v, f"{sm:.4f}", method_labels.get(method, method)])
        _add_table_from_list(doc, rows, ["Тип ТС", "Sm", "Метод"])

    # 4. ΣNp
    doc.add_page_break()
    _add_heading(doc, "4. Расчёт ΣNp", level=1)
    _add_formula_block(doc, "ΣNp = f_пол · Σ[N_m · Kc · Трдг · 0.7 · Sm_m · k_n] / 10^6",
                       "Формула 7.1 СТО 2.4-2013:")
    _add_formula_block(doc, "Kc = (q^Tсл - 1) / (q - 1)", "Коэффициент суммирования:")
    _add_table_from_dict(doc, {
        "ΣNp (проект), млн. авт.": f"{results.get('np_project', 0):.4f}",
        "ΣNp (факт), млн. авт.": f"{results.get('np_fact', 0):.4f}",
        "ΣNост (по нагрузке), млн. авт.": f"{results.get('residual', {}).get('np_residual', 0):.4f}",
        "ΣNост (по нагрузке), %": f"{results.get('residual', {}).get('percent_residual', 0):.2f}",
    })

    # 5. Прочность
    doc.add_page_break()
    _add_heading(doc, "5. Оценка прочности", level=1)
    kpr_tr = results.get("kpr_tr", 0)
    kpr_fact = results.get("kpr_fact", 0)
    if kpr_fact < 1.00:
        status = "ПРОЧНОСТЬ НЕДОСТАТОЧНА"
        rec = "Требуется капитальный ремонт или усиление."
    elif kpr_fact < kpr_tr:
        status = "Прочность ниже проектной, но работоспособна"
        rec = "Рекомендуется мониторинг."
    else:
        status = "Прочность соответствует проектной или выше"
        rec = "Конструкция работоспособна."
    _add_table_from_dict(doc, {
        "Kпр_тр": f"{kpr_tr:.3f}", "Kпр_факт": f"{kpr_fact:.3f}",
        "Статус": status, "Рекомендация": rec,
    })
    _add_heading(doc, "5.1. Ориентировочная оценка ресурса по прочности", level=2)
    _add_paragraph(doc, "ВНИМАНИЕ: методика НЕ утверждена нормативно.", italic=True)
    try:
        import calculations
        est = calculations.estimate_residual_by_kpr(kpr_fact, kpr_tr, results.get("np_project_for_kpr", 0.0))
        _add_table_from_dict(doc, {
            "Остаточный ресурс по прочности, %": f"{est['percent']:.2f}",
            "ΣNост (по прочности), млн. авт.": f"{est['np_residual']:.4f}",
        })
    except Exception as e:
        _add_paragraph(doc, f"Не удалось рассчитать: {e}", italic=True)

    # 6. Прогноз IRI
    if iri_forecast is not None:
        doc.add_page_break()
        _add_heading(doc, "6. Прогноз продольной ровности (IRI)", level=1)
        _add_formula_block(doc, "IRI_t = IRI_0 + a · (ΣN_ф.ед.)^b", "Формула 5.8 СТО 2.28-2016:")
        _add_table_from_dict(doc, {
            "Начальный IRI": f"{iri_forecast.get('iri_0', 0):.2f}",
            "Коэффициент a": f"{iri_forecast.get('coef_a', 0):.4f}",
            "Коэффициент b": f"{iri_forecast.get('coef_b', 0):.4f}",
        })
        years = iri_forecast.get("years", [])
        np_fed = iri_forecast.get("np_fed", [])
        iri = iri_forecast.get("iri", [])
        if years:
            rows = [[t,
                     f"{np_fed[i]:.4f}" if i < len(np_fed) else "—",
                     f"{iri[i]:.3f}" if i < len(iri) else "—"] for i, t in enumerate(years)]
            _add_table_from_list(doc, rows,
                ["Год", "ΣN_ф.ед., млн. авт.", "Прогноз IRI, м/км"])

    # 7. Итог
    doc.add_page_break()
    _add_heading(doc, "7. Итоговый вывод", level=1)
    percent_load = results.get("residual", {}).get("percent_residual", 0)
    _add_heading(doc, "7.1. Остаточный ресурс по нагрузке", level=2)
    _add_paragraph(doc, f"Остаточный ресурс: {percent_load:.1f}%")
    _add_heading(doc, "7.2. Прочность", level=2)
    _add_paragraph(doc, f"{status}. {rec}")
    _add_heading(doc, "7.3. Итоговая рекомендация", level=2)
    if kpr_fact < 1.00:
        _add_paragraph(doc, "КРИТИЧНО: необходим капитальный ремонт.", bold=True)
    elif percent_load < 25:
        _add_paragraph(doc, "Остаточный ресурс близок к исчерпанию.", bold=True)
    else:
        _add_paragraph(doc, "Конструкция работоспособна.", bold=True)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Отчёт сформирован автоматически в программе\n"
                    "«Расчёт остаточного ресурса дорожных конструкций»")
    run.italic = True
    run.font.size = Pt(10)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()